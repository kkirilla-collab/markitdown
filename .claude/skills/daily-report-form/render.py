# -*- coding: utf-8 -*-
"""Generic renderer for two-tier status reports (detailed + short/dynamic).

This module has NO knowledge of any specific project, company, or contract —
it only knows how to turn a plain data structure (see SKILL.md for the schema
and a worked example) into two formatted .docx files:

  - a detailed report: title block, KPI table, an arbitrary list of
    heading+table sections, a suggestions block, a footer note.
  - a short/dynamic report: title block, and an arbitrary list of
    heading+table blocks (meant for a KPI strip, a per-party status table,
    a day-over-day dynamics log, an open-correspondence tracker, etc).

Style is deliberately monochrome/"classic corporate" — black text on white,
black-fill/white-text table headers, no colour-coded (RAG) status highlighting.
Status is conveyed by the wording itself, rendered bold, never by hue — this
is a deliberate choice (not a limitation): plain black-and-white also holds up
when a report is printed or photocopied on a monochrome printer, unlike
colour-coding, which silently degrades to indistinguishable grey.

Callers build the `spec` dicts (usually Claude, after reading a new source
memo and extracting facts into this shape) and pass them to
`generate_detailed_report` / `generate_short_report`. Cells in a table row can
be a plain string, or a (text, opts) tuple produced by `status()`, `gap()` or
`blank()` below.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = "000000"       # all text — pure black, no theme colours
HEADER_FILL = "000000"  # table header row background
HEADER_FONT = "FFFFFF"  # table header row text (white-on-black)
MUTED = "595959"     # the one neutral grey used for de-emphasised/placeholder text


def fmt(n):
    """Format a number the RU corporate way: '20 659 286,79'."""
    s = f"{n:,.2f}"
    return s.replace(",", " ").replace(".", ",")


def status(text, level=None):
    """A table cell for a status/verdict value: bold black, no colour or fill.

    `level` (e.g. 'ok' | 'warn' | 'bad' | 'note') is accepted for the caller's
    own semantic bookkeeping (and so existing spec-building code doesn't need
    to change), but it no longer affects rendering — every status renders
    identically (bold black); the word itself carries the meaning.
    """
    return (text, {"bold": True})


def gap(text="не указано в справке"):
    """A table cell flagging a genuine gap in the source data (grey italic)."""
    return (text, {"italic": True, "color": MUTED})


def blank(text="—"):
    """A table cell that's an empty placeholder to be filled in later (grey italic)."""
    return (text, {"italic": True, "color": MUTED})


# --------------------------------------------------------------------------
# low-level docx helpers
# --------------------------------------------------------------------------

def _shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _set_cell(cell, text, bold=False, italic=False, color=None, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color or INK)
    return run


def _set_landscape(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.3)
    sec.bottom_margin = Cm(1.3)


def _add_table(doc, headers, rows, widths_cm=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        _set_cell(hdr_cells[i], h, bold=True, color=HEADER_FONT, size=9.5)
        _shade_cell(hdr_cells[i], HEADER_FILL)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if isinstance(val, tuple):
                text, opts = val
                _set_cell(cells[i], text, bold=opts.get("bold", False),
                          italic=opts.get("italic", False), color=opts.get("color"),
                          size=opts.get("size", 9.5))
            else:
                _set_cell(cells[i], str(val), size=9.5)
    if widths_cm:
        table.autofit = False
        for row in table.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def _add_bottom_border(paragraph, size=6, color=INK):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _h1(doc, text):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(INK)
    _add_bottom_border(p)


def _para(doc, text, bold=False, italic=False, size=10.5, color=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color or INK)


def _note(doc, text):
    _para(doc, text, italic=True, size=9.5, color=MUTED)


def _bullet(doc, lead, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if lead:
        r = p.add_run(lead)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor.from_string(INK)
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = RGBColor.from_string(INK)


def _set_base_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor.from_string(INK)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _render_table_block(doc, block):
    if block.get("intro"):
        for line in block["intro"]:
            _para(doc, line, size=10.5)
    t = block["table"]
    _add_table(doc, t["headers"], t["rows"], t.get("widths_cm"))
    for line in block.get("notes", []):
        _note(doc, line)


# --------------------------------------------------------------------------
# public API
# --------------------------------------------------------------------------

def generate_detailed_report(spec, out_path):
    """spec keys: title, subtitle_lines[], meta_line, top_notes[],
    sections[{heading, intro[], table{headers,rows,widths_cm}, notes[]}],
    suggestions{heading, intro, items[{lead,text}]}, footer_note.
    See SKILL.md for the full schema and a worked example."""
    doc = Document()
    _set_base_style(doc)
    _set_landscape(doc)

    title = doc.add_heading(spec["title"], level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor.from_string(INK)
        run.font.size = Pt(18)

    for line in spec.get("subtitle_lines", []):
        _para(doc, line, bold=True, size=11)
    for line in spec.get("top_notes", []):
        _note(doc, line)
    if spec.get("meta_line"):
        _para(doc, spec["meta_line"], size=10)

    for section in spec.get("sections", []):
        _h1(doc, section["heading"])
        _render_table_block(doc, section)

    sug = spec.get("suggestions")
    if sug:
        _h1(doc, sug["heading"])
        if sug.get("intro"):
            _para(doc, sug["intro"], size=10.5)
        for item in sug["items"]:
            _bullet(doc, item.get("lead", ""), item["text"])

    if spec.get("footer_note"):
        doc.add_page_break()
        _para(doc, spec["footer_note"], size=9.5, italic=True, color=MUTED)

    doc.save(out_path)
    return out_path


def generate_short_report(spec, out_path):
    """spec keys: title, subtitle_lines[], meta_line,
    blocks[{heading, intro[], table{headers,rows,widths_cm}, notes[]}], footer_note.
    See SKILL.md for the full schema and a worked example."""
    doc = Document()
    _set_base_style(doc)
    _set_landscape(doc)

    title = doc.add_heading(spec["title"], level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor.from_string(INK)
        run.font.size = Pt(16)

    for line in spec.get("subtitle_lines", []):
        _para(doc, line, size=10.5)
    if spec.get("meta_line"):
        _para(doc, spec["meta_line"], size=10)

    for block in spec.get("blocks", []):
        _h1(doc, block["heading"])
        _render_table_block(doc, block)

    if spec.get("footer_note"):
        _para(doc, spec["footer_note"], italic=True, size=9.5, color=MUTED)

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    import json
    import sys

    if len(sys.argv) != 4:
        print("usage: render.py <detailed|short> <spec.json> <out.docx>")
        raise SystemExit(1)
    kind, spec_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    fn = generate_detailed_report if kind == "detailed" else generate_short_report
    fn(spec, out_path)
    print(f"wrote {out_path}")
